"""Integration parsing tests."""

import tempfile
from pathlib import Path

from docx import Document

from integrations.documents import parse_docx_chapters
from integrations.sheets import parse_sheet_rows


class TestSheetParsing:
    def test_parse_rows_with_header(self):
        rows = [
            ["Number", "Title", "Content"],
            ["1", "Opening", "It was a dark night."],
            ["2", "Rising", "The hero awoke."],
        ]
        chapters = parse_sheet_rows(rows)
        assert len(chapters) == 2
        assert chapters[0]["number"] == 1
        assert chapters[0]["title"] == "Opening"

    def test_parse_rows_without_header(self):
        rows = [["1", "Start", "Hello world"]]
        chapters = parse_sheet_rows(rows)
        assert len(chapters) == 1
        assert chapters[0]["content"] == "Hello world"


class TestDocxParsing:
    def test_parse_chapter_headings(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "test.docx"
            doc = Document()
            doc.add_heading("Chapter 1: The Start", level=1)
            doc.add_paragraph("First paragraph.")
            doc.add_heading("Chapter 2: Continued", level=1)
            doc.add_paragraph("Second paragraph.")
            doc.save(str(path))

            chapters = parse_docx_chapters(str(path))
            assert len(chapters) == 2
            assert chapters[0]["title"] == "The Start"
            assert "First paragraph" in chapters[0]["content"]

    def test_single_document_becomes_chapter_one(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "test.docx"
            doc = Document()
            doc.add_paragraph("Just one block of text.")
            doc.save(str(path))

            chapters = parse_docx_chapters(str(path))
            assert len(chapters) == 1
            assert chapters[0]["number"] == 1
