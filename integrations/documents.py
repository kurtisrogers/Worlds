"""Document upload parsing service."""

from __future__ import annotations

import logging
import re

from django.utils import timezone
from docx import Document

from integrations.models import DocumentUpload
from integrations.sheets import sync_chapters_from_rows
from stories.models import Chapter

logger = logging.getLogger(__name__)

CHAPTER_HEADING = re.compile(r"^(chapter|ch\.?)\s+(\d+)[:\s]*(.*)$", re.IGNORECASE)


def parse_docx_chapters(file_path: str) -> list[dict]:
    """Extract chapters from a DOCX file based on heading patterns."""
    doc = Document(file_path)
    chapters: list[dict] = []
    current: dict | None = None
    content_lines: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            if current:
                content_lines.append("")
            continue

        match = CHAPTER_HEADING.match(text)
        if match:
            if current:
                current["content"] = "\n\n".join(content_lines).strip()
                chapters.append(current)
            number = int(match.group(2))
            title = match.group(3).strip() or f"Chapter {number}"
            current = {"number": number, "title": title, "content": ""}
            content_lines = []
        elif current:
            content_lines.append(text)

    if current:
        current["content"] = "\n\n".join(content_lines).strip()
        chapters.append(current)

    if not chapters:
        full_text = "\n\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())
        if full_text:
            chapters.append({"number": 1, "title": "Chapter 1", "content": full_text})

    return chapters


def process_document_upload(upload: DocumentUpload) -> int:
    """Parse an uploaded document and create/update chapters."""
    try:
        chapters_data = parse_docx_chapters(upload.file.path)
        count = sync_chapters_from_rows(upload.story, chapters_data)
        upload.status = DocumentUpload.Status.PARSED
        upload.chapters_created = count
        upload.parsed_at = timezone.now()
        upload.save()
        return count
    except Exception as exc:
        logger.exception("Document parse failed for upload %s", upload.pk)
        upload.status = DocumentUpload.Status.FAILED
        upload.error_message = str(exc)
        upload.save()
        return 0
